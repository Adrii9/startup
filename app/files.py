"""Storing uploaded files, and reading what is inside them.

Text is pulled out once, at upload, and kept beside the file. Doing it on every
read would make an assistant wait on a PDF parser, and doing it never would mean
the only way to use a document is to download it by hand.

Nothing here decides who may see a file; that is store.py's job.
"""

from __future__ import annotations

import hashlib
import mimetypes
import re
from pathlib import Path

# Room for a long report, with a ceiling so one document cannot fill the disk
# or a database row on its own.
MAX_TEXT = 400_000

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PROJECT_BYTES = 200 * 1024 * 1024

# Extensions we read as plain text whatever the browser called them. Source code
# arrives as application/octet-stream more often than not.
TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".env", ".log", ".sql", ".html", ".htm", ".css", ".xml",
    ".svg", ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".kt", ".c", ".h", ".cpp", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".swift", ".m", ".sh", ".bash", ".zsh", ".ps1",
    ".r", ".jl", ".lua", ".pl", ".dart", ".scala", ".tex", ".bib", ".gradle", ".dockerfile",
}

IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


def safe_name(name: str) -> str:
    """A display name with no directories and no surprises in it.

    Files are stored under a name we generate, so this is only what people and
    assistants see -- but a name with a path separator or a control character in
    it would still be a nasty thing to hand to anything downstream.
    """
    name = re.sub(r"[\x00-\x1f\x7f]", "", (name or "").replace("\\", "/").split("/")[-1])
    name = name.strip().strip(".") or "file"
    return name[:120]


def guess_mime(name: str, given: str = "") -> str:
    if given and given != "application/octet-stream":
        return given
    return mimetypes.guess_type(name)[0] or "application/octet-stream"


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def is_image(mime: str) -> bool:
    return mime in IMAGE_MIMES


def _from_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out = []
    for n, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append(f"--- page {n} ---\n{text}")
        if sum(len(x) for x in out) > MAX_TEXT:
            break
    return "\n\n".join(out)


def _from_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract(path: Path, name: str, mime: str) -> tuple[str, str]:
    """(text, state). state is what an assistant is told about this file:

    'text'   -- the content is right here
    'image'  -- no text, but read_file can hand over the picture itself
    'binary' -- stored and downloadable, nothing to read
    'empty'  -- we could read it and there was nothing in it (a scanned PDF)
    'error'  -- it should have been readable and was not
    """
    suffix = Path(name).suffix.lower()
    try:
        if is_image(mime):
            return "", "image"
        if mime == "application/pdf" or suffix == ".pdf":
            text = _from_pdf(path)
        elif suffix == ".docx" or mime.endswith("wordprocessingml.document"):
            text = _from_docx(path)
        elif mime.startswith("text/") or suffix in TEXT_SUFFIXES or mime in (
            "application/json", "application/xml", "application/javascript",
        ):
            text = path.read_bytes().decode("utf-8", errors="replace")
        else:
            return "", "binary"
    except Exception as e:                      # a corrupt PDF must not fail an upload
        return "", f"error: {type(e).__name__}"

    text = text.strip()
    if not text:
        return "", "empty"
    return text[:MAX_TEXT], "text"
